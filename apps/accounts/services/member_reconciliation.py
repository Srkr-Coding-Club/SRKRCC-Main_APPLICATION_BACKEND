"""Compare a past-member sheet (name, email, Club ID) with the accounts, row by row, and act on the differences.

Nothing is created or changed by reading a sheet. Changes happen only through the explicit actions below, one row (or one
outcome) at a time, all through the Club ID registry so each one is recorded with its reason.
"""
from __future__ import annotations

import csv
import io
import re
from difflib import SequenceMatcher

from django.conf import settings
from django.db import transaction
from django.db.models import Count
from django.db.models.functions import Lower, Upper
from django.utils import timezone
from django.utils.html import escape

from apps.accounts.models import (
    ClubIdEvent, ClubIdEventKind, ClubIdReservation, ReconciliationOutcome as Outcome, ReconciliationResolution as Resolution,
    ReconciliationRow, ReconciliationRun, ReservationStatus, User,
)
from apps.accounts.services import club_id_registry as registry
from apps.accounts.services.club_id_registry import ClubIdRegistryError
from apps.accounts.services.club_id_service import ClubIDService, InvalidClubIdError
from apps.accounts.services.member_import_service import MAX_ALLOWED_COLS, MAX_ALLOWED_ROWS, MAX_FILE_SIZE_BYTES, MemberImportService

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

EMAIL_PATTERN = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')
NAME_MATCH_RATIO = 0.72
MAX_NOTE = 300
FIELDS = ('name', 'email', 'club_id')
#: The outcomes an organizer has to look at; the rest are fine or are about the sheet itself.
NEEDS_DECISION = (
    Outcome.ID_MISSING_IN_APP, Outcome.ID_MISMATCH, Outcome.ID_HELD_BY_OTHER, Outcome.EMAIL_MISMATCH, Outcome.ID_RESERVED,
    Outcome.ID_RETIRED, Outcome.NO_ACCOUNT, Outcome.NO_ACCOUNT_NO_ID,
)
SIMPLE_CHOICES = {'KEEP_APP': Resolution.KEPT_APP, 'FOLLOW_UP': Resolution.FOLLOW_UP, 'DISMISS': Resolution.DISMISSED}


class ReconciliationError(Exception):
    def __init__(self, message: str, code: str, status: int = 400, **extra):
        super().__init__(message)
        self.message, self.code, self.status, self.extra = message, code, status, extra


# -- reading the sheet ---------------------------------------------------------------------------------------------------

def _cell(value) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _read_xlsx(file_obj, sheet: str | None) -> tuple[list[dict], str, list[str], list[list[str]]]:
    if openpyxl is None:
        raise ReconciliationError('Excel files cannot be read on this server. Upload a CSV instead.', 'NO_EXCEL')
    try:
        book = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    except Exception as exc:
        raise ReconciliationError('That file could not be read as an Excel workbook (.xlsx).', 'UNREADABLE') from exc
    sheets = [{'name': worksheet.title, 'rows': max((worksheet.max_row or 1) - 1, 0)} for worksheet in book.worksheets]
    if not sheets:
        raise ReconciliationError('The workbook has no sheets.', 'EMPTY')
    chosen = sheet or sheets[0]['name']
    if chosen not in {item['name'] for item in sheets}:
        raise ReconciliationError(f"There is no sheet called '{chosen}'.", 'NO_SUCH_SHEET')
    grid = [[_cell(value) for value in row] for row in book[chosen].iter_rows(values_only=True)]
    return sheets, chosen, *_split(grid)


def _read_csv(file_obj) -> tuple[list[dict], str, list[str], list[list[str]]]:
    raw = file_obj.read()
    text = raw.decode('utf-8-sig') if isinstance(raw, bytes) and _is_utf8(raw) else (raw.decode('latin-1') if isinstance(raw, bytes) else raw)
    grid = [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text))]
    return [], '', *_split(grid)


def _is_utf8(raw: bytes) -> bool:
    try:
        raw.decode('utf-8-sig')
        return True
    except UnicodeDecodeError:
        return False


def _split(grid: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    """(headers, rows below them). The first row with anything in it is the header; unnamed columns get a name."""
    grid = [row for row in grid if any(row)]
    if not grid:
        raise ReconciliationError('The sheet is empty.', 'EMPTY')
    width = min(max(len(row) for row in grid), MAX_ALLOWED_COLS)
    headers, seen = [], {}
    for index in range(width):
        name = (grid[0][index] if index < len(grid[0]) else '') or f'Column {index + 1}'
        seen[name] = seen.get(name, 0) + 1
        headers.append(name if seen[name] == 1 else f'{name} ({seen[name]})')
    return headers, grid[1:]


def read_sheet(file_obj, filename: str, sheet: str | None = None) -> dict:
    """The sheets, the headers of the chosen one, a suggested mapping, and every row as {header: text}."""
    if getattr(file_obj, 'size', 0) > MAX_FILE_SIZE_BYTES:
        raise ReconciliationError('The file is larger than 10 MB.', 'TOO_LARGE')
    lowered = (filename or '').lower()
    if lowered.endswith('.xlsx'):
        sheets, chosen, headers, body = _read_xlsx(file_obj, sheet)
    elif lowered.endswith('.csv'):
        sheets, chosen, headers, body = _read_csv(file_obj)
    else:
        raise ReconciliationError('Upload an Excel (.xlsx) or CSV (.csv) file. Older .xls files must be saved as .xlsx first.', 'UNSUPPORTED')
    if len(body) > MAX_ALLOWED_ROWS:
        raise ReconciliationError(f'The sheet has more than {MAX_ALLOWED_ROWS} rows. Split it and upload the parts.', 'TOO_MANY_ROWS')
    rows = [{header: (row[index] if index < len(row) else '') for index, header in enumerate(headers)} for row in body]
    return {'sheets': sheets, 'sheet': chosen, 'headers': headers, 'suggested': suggest_mapping(headers), 'rows': rows}


def suggest_mapping(headers: list[str]) -> dict[str, str]:
    """Which header looks like the name, the email and the Club ID ('' when none does)."""
    guessed = MemberImportService.auto_detect_mapping(headers)
    found = {field: '' for field in FIELDS}
    for header, target in guessed.items():
        if target == 'full_name':
            found['name'] = header
        elif target in ('email', 'club_id') and not found[target]:
            found[target] = header
    return found


def check_mapping(mapping: dict, headers: list[str]) -> dict[str, str]:
    clean = {field: (mapping or {}).get(field, '') or '' for field in FIELDS}
    for field, header in clean.items():
        if header and header not in headers:
            raise ReconciliationError(f"The sheet has no column '{header}'.", 'BAD_MAPPING')
    if not clean['email'] and not clean['club_id']:
        raise ReconciliationError('Choose the column that holds the email or the Club ID (or both).', 'BAD_MAPPING')
    return clean


# -- classifying a row ---------------------------------------------------------------------------------------------------

def _tokens(text: str) -> list[str]:
    return sorted(re.sub(r'[^a-z0-9 ]', ' ', (text or '').casefold()).split())


def names_agree(sheet_name: str, user: User) -> bool:
    """True when the names are plausibly the same person (order, case, punctuation and an extra initial do not matter)."""
    sheet_tokens, app_tokens = _tokens(sheet_name), _tokens(f'{user.first_name} {user.last_name}')
    if not sheet_tokens or not app_tokens:
        return True
    if set(sheet_tokens) <= set(app_tokens) or set(app_tokens) <= set(sheet_tokens):
        return True
    return SequenceMatcher(None, ' '.join(sheet_tokens), ' '.join(app_tokens)).ratio() >= NAME_MATCH_RATIO


def normalize_email(text: str) -> str:
    return (text or '').strip().lower()


def normalize_id(text: str) -> tuple[str, str]:
    """(canonical id or '', problem or '')."""
    text = (text or '').strip()
    if not text:
        return '', ''
    try:
        return ClubIDService.parse_club_id(text)['canonical_id'], ''
    except InvalidClubIdError:
        return '', f"'{text}' is not a Club ID like 25SCC277 or 21CCA041"


class _Lookup:
    """What the accounts, reservations and retired ids say about a set of emails and ids, fetched in a few queries."""

    def __init__(self, emails: set[str], ids: set[str]):
        self.users = {u.email.lower(): u for u in User.objects.annotate(lowered=Lower('email')).filter(lowered__in=emails)} if emails else {}
        self.holders = {u.club_id.upper(): u for u in User.objects.annotate(upper_id=Upper('club_id')).filter(upper_id__in=ids)} if ids else {}
        self.reserved = {r.club_id.upper(): r for r in ClubIdReservation.objects.filter(status=ReservationStatus.RESERVED)}
        self.retired = set(ClubIdEvent.objects.filter(kind=ClubIdEventKind.CLEARED).values_list('club_id', flat=True))


def classify(row: ReconciliationRow, lookup: _Lookup) -> tuple[str, str, User | None, User | None]:
    """(outcome, a sentence for the organizer, the account found by email, the account holding the sheet's id)."""
    email = normalize_email(row.sheet_email)
    club_id, id_problem = normalize_id(row.sheet_club_id)
    if id_problem:
        return Outcome.INVALID, f'{id_problem}.', None, None
    if email and not EMAIL_PATTERN.match(email):
        return Outcome.INVALID, f"'{row.sheet_email}' is not an email address.", None, None
    if not email and not club_id:
        return Outcome.INVALID, 'The row has neither an email nor a Club ID.', None, None

    user = lookup.users.get(email) if email else None
    holder = lookup.holders.get(club_id) if club_id else None

    if user and club_id:
        if user.club_id and user.club_id.upper() == club_id:
            if names_agree(row.sheet_name, user):
                return Outcome.MATCH, 'Email and Club ID agree with the account.', user, user
            return Outcome.NAME_DIFFERS, f"Email and Club ID agree, but the sheet says '{row.sheet_name}' and the account is '{user.first_name} {user.last_name}'.".replace("  ", " "), user, user
        if user.club_id:
            return Outcome.ID_MISMATCH, f'The account has Club ID {user.club_id}; the sheet says {club_id}.', user, holder
        if holder:
            return Outcome.ID_HELD_BY_OTHER, f'{club_id} belongs to {holder.email}, but the sheet gives it to {email}, who has no Club ID.', user, holder
        reservation = lookup.reserved.get(club_id)
        if reservation and reservation.reserved_for_email != email:
            return Outcome.ID_RESERVED, f'{club_id} is reserved for {reservation.reserved_for_email or "someone, held back by an organizer"}.', user, None
        if club_id in {item.upper() for item in lookup.retired}:
            return Outcome.ID_RETIRED, f'{club_id} was cleared from a member earlier and retired. An admin can still link it on purpose.', user, None
        return Outcome.ID_MISSING_IN_APP, f'The account exists and has no Club ID; {club_id} is free, so it can be linked.', user, None
    if user:
        text = f'The account has Club ID {user.club_id}.' if user.club_id else 'The account has no Club ID either.'
        return Outcome.NO_ID_IN_SHEET, f'The sheet has no Club ID for this email. {text}', user, None
    if holder:
        label = 'a different email' if email else 'no email in the sheet'
        return Outcome.EMAIL_MISMATCH, f'{club_id} is held by {holder.email} ({label}). It may be the same person with a new email.', None, holder
    if club_id:
        reservation = lookup.reserved.get(club_id)
        if reservation and reservation.reserved_for_email != email:
            return Outcome.ID_RESERVED, f'{club_id} is reserved for {reservation.reserved_for_email or "someone, held back by an organizer"}.', None, None
        if not email:
            return Outcome.NO_ACCOUNT, f'Nobody holds {club_id} and the sheet has no email, so nobody can be invited.', None, None
        already = ' It is already reserved for this email.' if reservation else ''
        return Outcome.NO_ACCOUNT, f'No account uses this email and nobody holds {club_id}.{already}', None, None
    return Outcome.NO_ACCOUNT_NO_ID, 'No account uses this email, and the sheet gives no Club ID.', None, None


def _lookup_for(rows) -> _Lookup:
    emails = {normalize_email(row.sheet_email) for row in rows if normalize_email(row.sheet_email)}
    ids = {normalize_id(row.sheet_club_id)[0] for row in rows if normalize_id(row.sheet_club_id)[0]}
    return _Lookup(emails, ids)


def refresh(row: ReconciliationRow, lookup: _Lookup | None = None, *, save: bool = True) -> ReconciliationRow:
    """Re-read what the accounts say now (after an action, or when time has passed). A repeated row stays a repeat."""
    if row.outcome == Outcome.DUPLICATE:
        return row
    outcome, detail, user, other = classify(row, lookup or _lookup_for([row]))
    row.outcome, row.detail, row.matched_user, row.other_user = outcome, detail[:500], user, other
    if save:
        row.save(update_fields=['outcome', 'detail', 'matched_user', 'other_user'])
    return row


# -- runs ----------------------------------------------------------------------------------------------------------------

@transaction.atomic
def create_run(actor, file_obj, filename: str, *, sheet: str | None, mapping: dict) -> ReconciliationRun:
    data = read_sheet(file_obj, filename, sheet)
    clean = check_mapping(mapping, data['headers'])
    run = ReconciliationRun.objects.create(source_filename=filename[:255], sheet_name=data['sheet'], mapping=clean, total_rows=len(data['rows']), created_by=actor)
    rows = []
    for index, source in enumerate(data['rows'], start=2):
        pick = lambda field: source.get(clean[field], '') if clean[field] else ''  # noqa: E731
        rows.append(ReconciliationRow(
            run=run, row_number=index, sheet_name=pick('name')[:200], sheet_email=pick('email')[:254], sheet_club_id=pick('club_id')[:60],
            raw={key: value[:300] for key, value in source.items()}, outcome=Outcome.INVALID,
        ))
    lookup = _lookup_for(rows)
    first_email: dict[str, int] = {}
    first_id: dict[str, int] = {}
    for row in rows:
        email, (club_id, _) = normalize_email(row.sheet_email), normalize_id(row.sheet_club_id)
        repeated = (first_email.get(email) if email else None) or (first_id.get(club_id) if club_id else None)
        if repeated:
            row.outcome, row.detail = Outcome.DUPLICATE, f'The same email or Club ID is already on row {repeated}.'
            continue
        if email:
            first_email[email] = row.row_number
        if club_id:
            first_id[club_id] = row.row_number
        refresh(row, lookup, save=False)
    ReconciliationRow.objects.bulk_create(rows, batch_size=500)
    return run


def summary(run: ReconciliationRun) -> dict:
    """Counts by outcome and by resolution, and how many rows still need a decision."""
    by_outcome = {key: 0 for key in Outcome.values}
    by_resolution = {key: 0 for key in Resolution.values}
    for row in run.rows.values('outcome', 'resolution').annotate(n=Count('pk')):
        by_outcome[row['outcome']] += row['n']
        by_resolution[row['resolution']] += row['n']
    open_rows = sum(run.rows.filter(outcome=key, resolution=Resolution.PENDING).count() for key in NEEDS_DECISION)
    return {
        'id': str(run.id), 'filename': run.source_filename, 'sheet': run.sheet_name, 'mapping': run.mapping, 'total': run.total_rows,
        'created_at': run.created_at, 'created_by': getattr(run.created_by, 'email', ''), 'by_outcome': by_outcome, 'by_resolution': by_resolution,
        'open': open_rows,
    }


def missing_from_sheet(run: ReconciliationRun):
    """Members in the app (holding a Club ID) who are not in this sheet, by neither email nor id."""
    emails = {normalize_email(e) for e in run.rows.exclude(sheet_email='').values_list('sheet_email', flat=True)}
    ids = {normalize_id(c)[0] for c in run.rows.exclude(sheet_club_id='').values_list('sheet_club_id', flat=True)} - {''}
    return (
        User.objects.exclude(club_id__isnull=True).exclude(club_id='').annotate(lowered=Lower('email'), upper_id=Upper('club_id'))
        .exclude(lowered__in=emails).exclude(upper_id__in=ids)
    )


# -- actions on one row --------------------------------------------------------------------------------------------------

def _reason(run: ReconciliationRun, row: ReconciliationRow, note: str = '') -> str:
    base = f'Past-member sheet {run.source_filename}, row {row.row_number}'
    return f'{base}: {note}' if note else base


def refresh_rows(rows) -> None:
    """Bring a batch of rows up to date with the accounts (one set of queries), saving only those that changed."""
    rows = [row for row in rows if row.outcome != Outcome.DUPLICATE]
    if not rows:
        return
    lookup, changed = _lookup_for(rows), []
    for row in rows:
        before = (row.outcome, row.detail, row.matched_user_id, row.other_user_id)
        refresh(row, lookup, save=False)
        if before != (row.outcome, row.detail, row.matched_user_id, row.other_user_id):
            changed.append(row)
    if changed:
        ReconciliationRow.objects.bulk_update(changed, ['outcome', 'detail', 'matched_user', 'other_user'], batch_size=500)


def recheck(run: ReconciliationRun) -> dict:
    """Look at every row again (after people registered, ids were linked elsewhere or reservations changed)."""
    before = {row.pk: row.outcome for row in run.rows.only('pk', 'outcome')}
    refresh_rows(list(run.rows.select_related('matched_user', 'other_user')))
    changed = sum(1 for row in run.rows.only('pk', 'outcome') if before[row.pk] != row.outcome)
    return {'changed': changed}


def _refusal_refreshes(action):
    """When an action is refused because the row is out of date, store the row's real state (outside the rolled-back transaction)."""
    def wrapper(actor, run, row_id, *args, **kwargs):
        try:
            return action(actor, run, row_id, *args, **kwargs)
        except ReconciliationError as exc:
            row = ReconciliationRow.objects.filter(run=run, pk=row_id).first()
            if row is not None and exc.code in ('NOT_APPLICABLE', 'CLUB_ID_UNAVAILABLE'):
                refresh(row)
            raise
    wrapper.__name__, wrapper.__doc__ = action.__name__, action.__doc__
    return wrapper


def _locked(run: ReconciliationRun, row_id) -> ReconciliationRow:
    row = ReconciliationRow.objects.select_for_update(of=('self',)).select_related('matched_user', 'other_user').filter(run=run, pk=row_id).first()
    if row is None:
        raise ReconciliationError('That row was not found.', 'NOT_FOUND', 404)
    return refresh(row)


def _done(actor, row: ReconciliationRow, resolution: str, note: str = '') -> ReconciliationRow:
    row.resolution, row.resolved_by, row.resolved_at, row.note = resolution, actor, timezone.now(), (note or '')[:MAX_NOTE]
    row.save(update_fields=['resolution', 'resolved_by', 'resolved_at', 'note'])
    return refresh(row)


def _need(row: ReconciliationRow, *allowed: str) -> None:
    if row.outcome not in allowed:
        raise ReconciliationError(f'This row is now "{row.get_outcome_display()}", so that does not apply. Reload the list.', 'NOT_APPLICABLE', 409, outcome=row.outcome)


@_refusal_refreshes
@transaction.atomic
def link(actor, run: ReconciliationRun, row_id, *, allow_retired: bool = False) -> ReconciliationRow:
    """Give the account found by email the Club ID the sheet has (only while the id is free)."""
    row = _locked(run, row_id)
    _need(row, Outcome.ID_MISSING_IN_APP, *( [Outcome.ID_RETIRED] if allow_retired else [] ))
    try:
        registry.assign(actor, row.matched_user_id, normalize_id(row.sheet_club_id)[0], reason=_reason(run, row), allow_retired=allow_retired)
    except ClubIdRegistryError as exc:
        raise ReconciliationError(exc.message, exc.code, exc.status) from exc
    return _done(actor, row, Resolution.LINKED)


@_refusal_refreshes
@transaction.atomic
def use_sheet_value(actor, run: ReconciliationRun, row_id, note: str) -> ReconciliationRow:
    """The sheet is right: the account's Club ID becomes the sheet's (the old id is freed, not retired)."""
    note = (note or '').strip()
    if not note:
        raise ReconciliationError('Say why the sheet is right: it is kept in the history.', 'REASON_REQUIRED')
    row = _locked(run, row_id)
    _need(row, Outcome.ID_MISMATCH)
    try:
        registry.replace(actor, row.matched_user_id, normalize_id(row.sheet_club_id)[0], reason=_reason(run, row, note))
    except ClubIdRegistryError as exc:
        raise ReconciliationError(exc.message, exc.code, exc.status) from exc
    return _done(actor, row, Resolution.USED_SHEET, note)


@_refusal_refreshes
@transaction.atomic
def decide(actor, run: ReconciliationRun, row_id, choice: str, note: str = '') -> ReconciliationRow:
    """Keep the app's value, mark the row for follow-up, or dismiss it. Nothing changes in the accounts."""
    if choice not in SIMPLE_CHOICES:
        raise ReconciliationError('Choose keep the app value, follow up or dismiss.', 'INVALID')
    row = _locked(run, row_id)
    return _done(actor, row, SIMPLE_CHOICES[choice], note)


def _invite_text(row: ReconciliationRow, club_id: str) -> tuple[str, str, str]:
    base = getattr(settings, 'FRONTEND_URL', 'http://localhost:3000').rstrip('/')
    url = f'{base}/signup'
    name = (row.sheet_name or row.sheet_email.split('@')[0]).strip()
    subject = 'SRKR Coding Club: create your account'
    text = (f'Hi {name},\n\nWe are bringing the SRKR Coding Club member records online and found you in our past-member list. '
            f'Your Club ID is {club_id}.\n\nCreate your account with this email address ({row.sheet_email}) and the ID is attached to it:\n{url}\n')
    html = (f'<p>Hi {escape(name)},</p><p>We are bringing the SRKR Coding Club member records online and found you in our past-member list. '
            f'Your Club ID is <strong>{escape(club_id)}</strong>.</p><p><a href="{escape(url)}">Create your account</a> with this email address '
            f'({escape(row.sheet_email)}) and the ID is attached to it.</p>')
    return subject, text, html


def send_invite_email(row_id: int) -> bool:
    """Runs on a background thread: looks the row up by id, then sends."""
    from apps.core.services.email_service import EmailNotificationService
    row = ReconciliationRow.objects.filter(pk=row_id).first()
    if row is None or not row.sheet_email:
        return False
    subject, text, html = _invite_text(row, normalize_id(row.sheet_club_id)[0])
    return EmailNotificationService.send_email(row.sheet_email, subject, text, html)


@_refusal_refreshes
@transaction.atomic
def invite(actor, run: ReconciliationRun, row_id) -> ReconciliationRow:
    """Reserve the sheet's Club ID for the email and send them a sign-up link. They get the id when they register."""
    row = _locked(run, row_id)
    _need(row, Outcome.NO_ACCOUNT)
    email, club_id = normalize_email(row.sheet_email), normalize_id(row.sheet_club_id)[0]
    if not email or not club_id:
        raise ReconciliationError('An invitation needs both an email and a Club ID.', 'NOT_APPLICABLE', 409)
    if row.invited_at:
        raise ReconciliationError('This person was already invited.', 'ALREADY_INVITED', 409)
    existing = ClubIdReservation.objects.filter(club_id__iexact=club_id, reserved_for_email=email, status=ReservationStatus.RESERVED).first()
    if existing is None:
        try:
            registry.reserve(actor, email=email, club_id=club_id, reason=_reason(run, row), source='RECONCILIATION')
        except ClubIdRegistryError as exc:
            raise ReconciliationError(exc.message, exc.code, exc.status) from exc
    row.invited_at = timezone.now()
    row.save(update_fields=['invited_at'])
    from apps.core.tasks import run_in_background
    row_pk = row.pk
    transaction.on_commit(lambda: run_in_background(lambda: send_invite_email(row_pk)))
    return _done(actor, row, Resolution.INVITED)


def bulk(actor, run: ReconciliationRun, action: str, outcome: str, expected: int) -> dict:
    """The same action on every open row with that outcome. `expected` is the count the person saw, so a changed list is not acted on."""
    allowed = {'link': Outcome.ID_MISSING_IN_APP, 'invite': Outcome.NO_ACCOUNT}
    if action not in allowed or allowed[action] != outcome:
        raise ReconciliationError('Choose link (accounts without an id) or invite (no account).', 'INVALID')
    rows = list(run.rows.filter(outcome=outcome, resolution=Resolution.PENDING).order_by('row_number'))
    if len(rows) != expected:
        raise ReconciliationError(f'The list changed: {len(rows)} rows are open now, not {expected}. Reload and check again.', 'STALE', 409, open=len(rows))
    done, failed = 0, []
    for row in rows:
        try:
            (link if action == 'link' else invite)(actor, run, row.pk)
            done += 1
        except (ReconciliationError, ClubIdRegistryError) as exc:
            failed.append({'row_number': row.row_number, 'error': getattr(exc, 'message', str(exc))})
    return {'done': done, 'failed': failed}
